"""doyz -> Word (.docx)。

图表两种模式：
  native（默认）：Word 原生图表，双击可改数据、可在 Word 里换类型
  image        ：图表渲染为静态图片（兼容性最好，版式不会跑偏）
"""

from __future__ import annotations

import base64
import os
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import chart_img, model


# ---------------------------------------------------------------- 基础工具

def _rgb(hex_str, default="000000"):
    if not hex_str:
        return RGBColor.from_string(default)
    return RGBColor.from_string(str(hex_str).lstrip("#"))


def _set_font(run, name, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = _rgb(color)


def _style_base(doc, style: dict, kind: str):
    font = style.get("font", "微软雅黑")
    hfont = style.get("headingFont") or font
    size = float(style.get("fontSizePt") or 11)
    accent = style.get("accent", "C2410C")

    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)

    sizes = {1: 18, 2: 15, 3: 13, 4: 12, 5: 11, 6: 11}
    for lvl in range(1, 7):
        try:
            st = doc.styles["Heading %d" % lvl]
        except KeyError:
            continue
        st.font.name = hfont
        st.element.rPr.rFonts.set(qn("w:eastAsia"), hfont)
        st.font.size = Pt(sizes[lvl])
        st.font.bold = True
        st.font.color.rgb = _rgb(accent if lvl <= 2 else "1F2937")
    return font, hfont, size, accent


def _page_setup(doc, page: dict):
    w, h = model.PAGE_SIZES_CM.get(page.get("size", "A4"), (21.0, 29.7))
    sec = doc.sections[0]
    if page.get("orientation") == "landscape":
        w, h = max(w, h), min(w, h)
        sec.orientation = WD_ORIENT.LANDSCAPE
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width = Cm(w)
    sec.page_height = Cm(h)
    m = page.get("marginCm") or [2.54] * 4
    sec.top_margin, sec.right_margin = Cm(m[0]), Cm(m[1])
    sec.bottom_margin, sec.left_margin = Cm(m[2]), Cm(m[3])
    return Cm(w - m[1] - m[3])


def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), str(hex_color).lstrip("#"))
    tcPr.append(shd)


def _add_page_number(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    run._r.addnext(fld)


def _resolve_image(src: str, base_dir: str) -> str:
    if os.path.exists(src):
        return src
    p = os.path.join(base_dir, src)
    if os.path.exists(p):
        return p
    if src.startswith("data:"):
        head, b64 = src.split(",", 1)
        ext = ".png"
        if "jpeg" in head or "jpg" in head:
            ext = ".jpg"
        elif "svg" in head:
            ext = ".svg"
        fd, path = tempfile.mkstemp(suffix=ext)
        with os.fdopen(fd, "wb") as f:
            f.write(base64.b64decode(b64))
        return path
    raise model.DoyzError("图片资源找不到: %s" % src)


# ---------------------------------------------------------------- 图表

def _add_native_chart(doc, chart, chart_type_enum, style: dict, content_width):
    """python-docx 没有图表 API：先埋标记，保存后由 word_chart 注入原生图表。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("@@DOYZCHART:%s@@" % chart.get("id"))
    run.font.size = Pt(1)
    return p


def _add_image_chart(doc, chart, style: dict, tmpdir: str):
    opt = chart.get("options") or {}
    png = os.path.join(tmpdir, "%s.png" % chart.get("id", "chart"))
    chart_img.render(chart, png, style)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width = Cm(float(opt.get("widthCm") or 15))
    p.add_run().add_picture(png, width=width)
    return p


# ---------------------------------------------------------------- 块渲染

def _render_runs(p, block, font, size, accent):
    runs = block.get("runs")
    text = block.get("text", "")
    if runs:
        for r in runs:
            run = p.add_run(str(r.get("t", "")))
            _set_font(run, font, size, r.get("b"), r.get("c"), r.get("i"))
    else:
        run = p.add_run(str(text))
        _set_font(run, font, size, block.get("bold"), block.get("color"), block.get("italic"))
    return p


def _render_block(doc, block, ctx):
    font, hfont, size, accent = ctx["font"], ctx["hfont"], ctx["size"], ctx["accent"]
    style = ctx["style"]
    btype = block.get("type", "paragraph")

    if btype == "heading":
        p = doc.add_heading(level=int(block.get("level") or 1))
        p.text = ""
        _render_runs(p, block, hfont, None, accent)
        return

    if btype == "paragraph":
        p = doc.add_paragraph()
        align = block.get("align")
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align == "justify":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _render_runs(p, block, font, size, accent)
        p.paragraph_format.space_after = Pt(6)
        return

    if btype == "bullets":
        st = "List Number" if block.get("ordered") else "List Bullet"
        for it in block.get("items") or []:
            p = doc.add_paragraph(style=st)
            _render_runs(p, {"text": it}, font, size, accent)
        return

    if btype == "quote":
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        run = p.add_run(str(block.get("text", "")))
        _set_font(run, font, size, None, "6B7280")
        return

    if btype == "code":
        p = doc.add_paragraph()
        run = p.add_run(str(block.get("text", "")))
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9.5)
        return

    if btype == "pagebreak":
        doc.add_page_break()
        return

    if btype == "table":
        header = block.get("header") or []
        rows = block.get("rows") or []
        cols = max([len(header)] + [len(r) for r in rows]) if (header or rows) else 0
        if not cols:
            return
        t = doc.add_table(rows=1, cols=cols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if header:
            for i, h in enumerate(header):
                cell = t.rows[0].cells[i]
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(h))
                _set_font(run, font, size - 0.5, True)
                _shade(cell, accent)
                run.font.color.rgb = RGBColor.from_string("FFFFFF")
        for r in rows:
            cells = t.add_row().cells
            for i in range(cols):
                v = r[i] if i < len(r) else ""
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                _set_font(run, font, size - 0.5)
        if block.get("caption"):
            _caption(doc, block["caption"], font, size)
        return

    if btype == "image":
        src = _resolve_image(block.get("src", ""), ctx["base_dir"])
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(src, width=Cm(float(block.get("widthCm") or 14)))
        if block.get("caption"):
            _caption(doc, block["caption"], font, size)
        return

    if btype == "chart":
        chart = model.get_chart(ctx["doc"], block.get("ref") or block.get("chart"))
        _caption_or_title(doc, chart.get("title"), font, size, accent, as_title=True)
        if ctx["chart_mode"] == "image":
            _add_image_chart(doc, chart, style, ctx["tmpdir"])
        else:
            _add_native_chart(doc, chart, None, style, ctx["content_width"])
        if block.get("caption"):
            _caption(doc, block["caption"], font, size)
        return

    if btype == "caption":
        _caption(doc, block.get("text", ""), font, size)
        return

    raise model.DoyzError("未知的块类型: %s" % btype)


def _caption_or_title(doc, text, font, size, accent, as_title=False):
    if not text:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    _set_font(run, font, size + (1 if as_title else -1.5), as_title, accent if as_title else "6B7280")
    p.paragraph_format.space_after = Pt(4)


def _caption(doc, text, font, size):
    _caption_or_title(doc, text, font, size, None, as_title=False)


# ---------------------------------------------------------------- 入口

def export(doc_data: dict, out_path: str, chart_mode: str = "native") -> str:
    if doc_data.get("kind") != "document":
        raise model.DoyzError("Word 导出要求 kind=document，当前为 %s" % doc_data.get("kind"))
    if chart_mode not in ("native", "image"):
        raise model.DoyzError("chart_mode 只能是 native 或 image")

    style = dict(model.DEFAULT_STYLE)
    style.update(doc_data.get("styles") or {})
    page = dict(model.DEFAULT_PAGE)
    page.update(doc_data.get("page") or {})

    doc = Document()
    content_width = _page_setup(doc, page)
    font, hfont, size, accent = _style_base(doc, style, "document")

    meta = doc_data.get("meta") or {}
    if meta.get("title"):
        sec = doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.text = ""
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run(str(meta["title"]))
        _set_font(run, font, 9, None, "9CA3AF")
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number(fp)
        core = doc.core_properties
        core.title = str(meta.get("title") or "")
        core.author = str(meta.get("author") or "")
        core.comments = "由 doyz 导出（图表模式: %s）" % chart_mode

    ctx = {
        "doc": doc_data, "style": style, "font": font, "hfont": hfont,
        "size": size, "accent": accent, "chart_mode": chart_mode,
        "content_width": content_width,
        "base_dir": doc_data.get("_base_dir") or os.path.dirname(os.path.abspath(out_path)),
        "tmpdir": doc_data.get("_tmpdir") or tempfile.gettempdir(),
    }
    for block in doc_data.get("blocks") or []:
        _render_block(doc, block, ctx)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or ".", exist_ok=True)
    doc.save(out_path)

    if chart_mode == "native":
        from . import word_chart
        word_chart.inject(out_path, doc_data, style)
    return out_path
